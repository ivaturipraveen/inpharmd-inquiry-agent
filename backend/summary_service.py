"""LLM-based answer extraction from raw call transcripts.

If the voice agent forgot to call `submit_answer` (or returned no summary),
we can extract a clean clinical answer from the full transcript using OpenAI.

Requires OPENAI_API_KEY env var. Falls back gracefully if not configured.
"""
from __future__ import annotations

import logging
import os
from typing import Optional

log = logging.getLogger("inquiry.summary")


class SummaryConfigError(RuntimeError):
    """Raised when OpenAI is not configured or the call fails."""


_client = None


def _get_client():
    global _client
    if _client is not None:
        return _client
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise SummaryConfigError(
            "OPENAI_API_KEY not set. Add it to backend/.env to enable AI extraction."
        )
    try:
        from openai import OpenAI
    except ImportError:
        raise SummaryConfigError(
            "openai package not installed. Run: pip install openai>=1.30"
        )
    _client = OpenAI(api_key=api_key)
    return _client


_SYSTEM = (
    "You are extracting the clinical answer from a medical-information phone call transcript. "
    "The call was between an InpharmD inquiry agent (caller) and a pharmaceutical manufacturer's "
    "Medical Information desk (representative). Read the transcript and extract ONLY the clinical "
    "answer the rep provided. Be faithful — do not interpret, do not add caveats unless the rep said "
    "them. If the rep cited a reference (case ID, package insert section, document name, case number), "
    "include it. If no clinical answer was given (voicemail, wrong number, rep declined), say so plainly "
    "in one short sentence."
)


_MODEL = "gpt-4o-mini"


def extract_answer_from_transcript(
    *,
    question: str,
    manufacturer: str,
    transcript: str,
) -> str:
    """Use OpenAI to extract a clean clinical answer from a call transcript."""
    client = _get_client()

    user = (
        f"MANUFACTURER: {manufacturer}\n"
        f"PHARMACIST'S QUESTION: {question}\n\n"
        f"TRANSCRIPT:\n{transcript}\n\n"
        "Extract the rep's clinical answer in 1–3 sentences. Include any references they "
        "cited. If no answer was given, say so plainly."
    )

    try:
        resp = client.chat.completions.create(
            model=_MODEL,
            messages=[
                {"role": "system", "content": _SYSTEM},
                {"role": "user", "content": user},
            ],
            max_tokens=400,
            temperature=0.2,
        )
        answer = (resp.choices[0].message.content or "").strip()
        log.info("Extracted answer (%d chars) via %s", len(answer), _MODEL)
        return answer
    except Exception as e:
        log.exception("OpenAI extraction failed: %s", e)
        raise SummaryConfigError(f"OpenAI call failed: {e}")


_EMAIL_SYSTEM = (
    "You are extracting the clinical answer from a medical-information email reply. "
    "The reply is from a pharmaceutical manufacturer's Medical Information desk, answering "
    "a pharmacist's clinical question that InpharmD forwarded. Read the reply and extract ONLY "
    "the clinical answer the manufacturer provided. Be faithful — do not interpret, do not add "
    "caveats unless they wrote them. If they cited a reference (case ID, package insert section, "
    "document name, case number), include it. Ignore email signatures, disclaimers, legal "
    "boilerplate, and quoted text from earlier messages. If no clinical answer was given "
    "(auto-reply, acknowledgement only, request for more info, declined), say so plainly in one short sentence."
)


def extract_answer_from_email(
    *,
    question: str,
    manufacturer: str,
    reply_text: str,
) -> str:
    """Use OpenAI to extract a clean clinical answer from an email reply body."""
    client = _get_client()

    user = (
        f"MANUFACTURER: {manufacturer}\n"
        f"PHARMACIST'S QUESTION: {question}\n\n"
        f"MANUFACTURER'S EMAIL REPLY:\n{reply_text}\n\n"
        "Extract their clinical answer in 1–3 sentences. Include any references they "
        "cited. If no answer was given, say so plainly."
    )

    try:
        resp = client.chat.completions.create(
            model=_MODEL,
            messages=[
                {"role": "system", "content": _EMAIL_SYSTEM},
                {"role": "user", "content": user},
            ],
            max_tokens=400,
            temperature=0.2,
        )
        answer = (resp.choices[0].message.content or "").strip()
        log.info("Extracted email answer (%d chars) via %s", len(answer), _MODEL)
        return answer
    except Exception as e:
        log.exception("OpenAI email extraction failed: %s", e)
        raise SummaryConfigError(f"OpenAI call failed: {e}")


_SIG_STRIP_SYSTEM = (
    "You receive the plain-text body of a manufacturer's email reply. Your ONLY job is to "
    "return the same body with the email signature, disclaimer, and confidentiality notice "
    "removed. Preserve the answer text VERBATIM — do not paraphrase, summarize, reword, "
    "fix grammar, add caveats, or change formatting in any way. Do not add a greeting, a "
    "closing, or any extra text. Return only the cleaned body. If the body has no signature "
    "to remove, return it exactly as received."
)


def strip_signature_with_ai(reply_text: str) -> str:
    """Use OpenAI to strip an email signature when regex-based stripping leaves
    one behind. The model is instructed to preserve the answer text verbatim
    and only remove the signature/disclaimer block. Falls back to the input
    on any error so we never lose the reply."""
    text = (reply_text or "").strip()
    if not text:
        return text
    try:
        client = _get_client()
    except SummaryConfigError:
        return text

    try:
        resp = client.chat.completions.create(
            model=_MODEL,
            messages=[
                {"role": "system", "content": _SIG_STRIP_SYSTEM},
                {"role": "user", "content": text},
            ],
            max_tokens=1200,
            temperature=0.0,
        )
        cleaned = (resp.choices[0].message.content or "").strip()
        if not cleaned:
            return text
        # Safety: if the model returned something dramatically shorter than the
        # input, it probably over-stripped — keep the manual version instead.
        if len(cleaned) < max(40, int(len(text) * 0.20)):
            log.warning("AI signature strip returned too little; keeping manual version")
            return text
        log.info("AI signature strip: %d → %d chars", len(text), len(cleaned))
        return cleaned
    except Exception as e:
        log.warning("AI signature strip failed (%s); keeping manual version", e)
        return text


def is_configured() -> bool:
    return bool(os.getenv("OPENAI_API_KEY"))


_PDF_SUMMARY_SYSTEM = (
    "You are summarizing a manufacturer-supplied medical-information PDF for a "
    "pharmacist. Read the attached document text and write ONE concise, clinically-"
    "useful answer to the inquiry. Be specific: include doses, durations, stability "
    "data, study references, or other concrete details mentioned in the document. "
    "Do not invent anything that is not in the text. If the document does not "
    "directly answer the question, say so and summarize what it DOES contain. "
    "Return plain prose — no preamble like 'Sure' or 'Here is'."
)


def summarize_pdf(question: str, manufacturer: str, pdf_text: str) -> str:
    """Summarize PDF body text into an answer to the inquiry. Caller handles
    'not configured' by checking is_configured() up front."""
    if not is_configured():
        raise SummaryConfigError("OPENAI_API_KEY not set")
    client = _get_client()
    # Cap input — these PDFs can be huge, and we don't need every page to write
    # a useful summary. ~24k chars ≈ 6k tokens, comfortable for any current model.
    snippet = (pdf_text or "").strip()
    if len(snippet) > 24_000:
        snippet = snippet[:24_000] + "\n…[truncated]"
    user = (
        f"Inquiry sent to {manufacturer}:\n\n{question.strip()}\n\n"
        f"PDF the manufacturer replied with:\n\n{snippet}"
    )
    try:
        resp = client.chat.completions.create(
            model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            temperature=0.2,
            messages=[
                {"role": "system", "content": _PDF_SUMMARY_SYSTEM},
                {"role": "user", "content": user},
            ],
        )
        out = (resp.choices[0].message.content or "").strip()
        log.info("PDF summary produced: %d chars from %d chars", len(out), len(snippet))
        return out
    except Exception as e:
        log.warning("PDF summary failed: %s", e)
        raise


def extract_pdf_text(pdf_bytes: bytes) -> str:
    """Extract plain text from a PDF using pypdf. Returns "" on failure."""
    try:
        from pypdf import PdfReader
        from io import BytesIO
        reader = PdfReader(BytesIO(pdf_bytes))
        chunks = []
        for page in reader.pages:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n\n".join(c for c in chunks if c.strip())
    except Exception as e:
        log.warning("PDF text extract failed: %s", e)
        return ""


def extract_document_text(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from a document based on its file extension.

    Supports: .pdf, .docx, .doc, .csv, .xlsx, .xls
    Returns "" on failure or unsupported type.
    """
    ext = (filename or "").rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        return extract_pdf_text(file_bytes)

    if ext == "docx":
        try:
            import docx  # python-docx
            from io import BytesIO
            doc = docx.Document(BytesIO(file_bytes))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            log.warning("DOCX text extract failed: %s", e)
            return ""

    if ext == "doc":
        # Legacy binary Word format — attempt a best-effort text extraction by
        # decoding bytes and stripping non-printable characters.
        try:
            text = file_bytes.decode("latin-1", errors="ignore")
            import re
            text = re.sub(r"[^\x20-\x7e\n\r\t]", " ", text)
            text = re.sub(r" {4,}", " ", text)
            return text.strip()
        except Exception as e:
            log.warning("DOC text extract failed: %s", e)
            return ""

    if ext == "csv":
        try:
            import csv
            from io import StringIO
            decoded = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(StringIO(decoded))
            rows = ["\t".join(row) for row in reader]
            return "\n".join(rows)
        except Exception as e:
            log.warning("CSV text extract failed: %s", e)
            return ""

    if ext == "xlsx":
        try:
            import openpyxl
            from io import BytesIO
            wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            chunks = []
            for sheet in wb.worksheets:
                chunks.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join("" if v is None else str(v) for v in row)
                    if line.strip():
                        chunks.append(line)
            return "\n".join(chunks)
        except Exception as e:
            log.warning("XLSX text extract failed: %s", e)
            return ""

    if ext == "xls":
        try:
            import xlrd
            from io import BytesIO
            wb = xlrd.open_workbook(file_contents=file_bytes)
            chunks = []
            for sheet in wb.sheets():
                chunks.append(f"[Sheet: {sheet.name}]")
                for rx in range(sheet.nrows):
                    line = "\t".join(str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols))
                    if line.strip():
                        chunks.append(line)
            return "\n".join(chunks)
        except Exception as e:
            log.warning("XLS text extract failed: %s", e)
            return ""

    log.info("No text extractor for file extension '%s'", ext)
    return ""
